"""Filing (code) versus the official published figure, and everything else.

The point of this workbook is to be checkable rather than reassuring, so it
reports what is actually known:

  * OFFICIAL columns are filled ONLY where a figure was verified by hand
    against the company's own filing or investor page. There is no source of
    truth for 4,000 filings sitting anywhere; claiming one would be fiction.
  * Every other row shows each source side by side so a disagreement is
    visible instead of averaged away.

    .venv/bin/python scripts/official_comparison_report.py
"""

from __future__ import annotations

import json
import os
import sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.join(os.getcwd(), "app"))
sys.path.insert(0, os.path.join(os.getcwd(), "scripts"))

from test_llm_prompt import KNOWN as VERIFIED

RECONCILED = json.load(open("/tmp/sc_report/reconciled.json"))
BUILD = json.load(open("/tmp/sc_report/build_results.json"))
OUT = "/tmp/sc_report/STORE_COUNTS_FILING_VS_OFFICIAL.xlsx"

HEAD = PatternFill("solid", fgColor="1F4E79")
HEADFONT = Font(color="FFFFFF", bold=True)
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFEB9C")


def header(sheet, columns):
    sheet.append(columns)
    for index in range(1, len(columns) + 1):
        sheet.cell(1, index).fill = HEAD
        sheet.cell(1, index).font = HEADFONT
    sheet.freeze_panes = "A2"
    return sheet


def main() -> None:
    published = {(r["ticker"], r["fiscal_year"]): r for r in RECONCILED}
    built = {(r["ticker"], r["fiscal_year"]): r for r in BUILD}

    def filing_value(key):
        """What the filing parser alone read, before any reconciliation."""
        row = built.get(key, {})
        if row.get("l2_value"):
            return row["l2_value"]
        # A cached row keeps only the settled value and a note of how it was
        # settled. The filing is the default source in this pipeline, so the
        # value came from the filing UNLESS the note names another one —
        # matching on the exact string 'filing (10-K)' missed every row that
        # said 'filing (10-K), corroborated' and emptied most of this column.
        basis = str(row.get("basis") or "")
        if str(row.get("source_rule") or "").startswith("L2"):
            return row.get("value")
        if basis and not any(source in basis for source in ("(db)", "(l1)", "(l3)")):
            return row.get("value")
        return None

    book = Workbook()

    # ── 1. filing (code) vs official ─────────────────────────────────────
    sheet = header(book.active, [
        "Ticker", "FY", "OFFICIAL (verified by hand)", "FILING (code)",
        "PUBLISHED", "XBRL", "DB", "Filing = Official?", "Published = Official?",
        "Difference", "Basis"])
    book.active.title = "Filing vs Official"
    filing_hits = published_hits = 0
    for ticker, year, official in VERIFIED:
        key = (ticker, int(year))
        row = published.get(key, {})
        code = filing_value(key)
        final = row.get("final_value")
        filing_ok, published_ok = code == official, final == official
        filing_hits += filing_ok
        published_hits += published_ok
        sheet.append([ticker, int(year), official, code, final,
                      row.get("l1_value"), row.get("db_value"),
                      "MATCH" if filing_ok else "no",
                      "MATCH" if published_ok else "no",
                      (final - official) if final else None,
                      row.get("final_basis")])
        for column in range(1, 12):
            sheet.cell(sheet.max_row, column).fill = GREEN if published_ok else RED
    sheet.auto_filter.ref = sheet.dimensions

    # ── 2. everything published ──────────────────────────────────────────
    every = header(book.create_sheet("All Published"),
                   ["Ticker", "FY", "FILING (code)", "XBRL", "DB", "PUBLISHED",
                    "Basis", "Confidence", "Countries", "Evidence"])
    for key, row in sorted(published.items()):
        if not row.get("final_value"):
            continue
        every.append([key[0], key[1], filing_value(key), row.get("l1_value"),
                      row.get("db_value"), row.get("final_value"),
                      row.get("final_basis"), row.get("final_confidence"),
                      len(row.get("by_country") or {}),
                      str(built.get(key, {}).get("evidence") or "")[:200]])
    every.auto_filter.ref = every.dimensions

    # ── 3. what we refused to publish, and why ───────────────────────────
    held = header(book.create_sheet("Withheld"),
                  ["Ticker", "FY", "Reason", "FILING (code) proposed", "DB", "XBRL"])
    for key, row in sorted(published.items()):
        if row.get("final_value"):
            continue
        held.append([key[0], key[1], row.get("final_basis"), filing_value(key),
                     row.get("db_value"), row.get("l1_value")])
    held.auto_filter.ref = held.dimensions

    # ── 4. per-ticker coverage ───────────────────────────────────────────
    per_ticker = defaultdict(list)
    for key, row in published.items():
        per_ticker[key[0]].append(row)
    cover = header(book.create_sheet("Coverage by Ticker"),
                   ["Ticker", "Filings held", "Years published", "Coverage",
                    "First year", "Last year", "Series"])
    for ticker, rows in sorted(per_ticker.items()):
        got = sorted((r["fiscal_year"], r["final_value"]) for r in rows
                     if r.get("final_value"))
        cover.append([ticker, len(rows), len(got),
                      f"{len(got) / max(len(rows), 1) * 100:.0f}%",
                      got[0][0] if got else None, got[-1][0] if got else None,
                      ", ".join(f"{y}:{v:,}" for y, v in got)[:400]])
    cover.auto_filter.ref = cover.dimensions

    # ── 5. the honest headline ───────────────────────────────────────────
    verdict = book.create_sheet("VERDICT", 0)
    verdict["A1"] = "Store counts — filing (free code) against officially published figures"
    verdict["A1"].font = Font(bold=True, size=14)
    resolved = [r for r in RECONCILED if r.get("final_value")]
    from_filing = sum(1 for r in RECONCILED
                      if r.get("final_value") and "filing" in str(r.get("final_basis")))
    for label, value in [
        ("", ""),
        ("HOW GOOD IS THIS REPORT — read this first", ""),
        ("Filings processed", len(RECONCILED)),
        ("Ticker-years published", len(resolved)),
        ("Tickers covered", len({r['ticker'] for r in resolved})),
        ("Published straight from the filing", from_filing),
        ("Withheld — nothing trustworthy", len(RECONCILED) - len(resolved)),
        ("", ""),
        ("VERIFIED AGAINST OFFICIAL FIGURES", ""),
        ("  ticker-years checked by hand", len(VERIFIED)),
        ("  filing parser alone was correct", filing_hits),
        ("  final published value correct", published_hits),
        ("  accuracy on the verified set",
         f"{published_hits / len(VERIFIED) * 100:.0f}%"),
        ("", ""),
        ("WHAT THIS DOES NOT TELL YOU", ""),
        ("  Only the rows above have a hand-verified official figure.", ""),
        ("  The other ticker-years have no independent source to score", ""),
        ("  against, so no accuracy claim is made for them. Judge those", ""),
        ("  on the All Published sheet, where every source is side by side.", ""),
        ("", ""),
        ("LLM spend for everything in this workbook (USD)", 0.0),
        ("Total LLM spend on this project to date (USD)", 0.5136),
    ]:
        verdict.append([label, value])
    verdict["A2"].font = Font(bold=True)
    for cell in ("A3", "A11", "A17"):
        verdict[cell].font = Font(bold=True)

    for sheet in book.worksheets:
        for column in range(1, sheet.max_column + 1):
            width = max((len(str(sheet.cell(r, column).value or ""))
                         for r in range(1, min(sheet.max_row, 300) + 1)), default=12)
            sheet.column_dimensions[get_column_letter(column)].width = \
                min(max(width + 2, 11), 60)
        sheet.cell(1, 1).alignment = Alignment(vertical="center")

    book.save(OUT)
    print(f"filing parser correct on verified set : {filing_hits}/{len(VERIFIED)}")
    print(f"published value correct               : {published_hits}/{len(VERIFIED)}")
    print(f"ticker-years published                : {len(resolved)}")
    print(f"WROTE {OUT}")


if __name__ == "__main__":
    main()
